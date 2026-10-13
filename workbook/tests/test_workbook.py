import datetime as dt

import openpyxl
import pytest

from scoring import perp_fields
from workbook import fill_perps
from workbook.sheet import Sheet, SheetError, open_book

UTC = dt.timezone.utc
D = dt.date(2026, 10, 13)
CFG = {"map": {"GOLD": "KXGOLDPERP", "OIL": None},
       "fields": {"perp_funding_z": {"window_days": 30, "min_settlements": 15, "clip": 5},
                  "perp_oi_chg_pct": {"lookback_hours": 24}}}


def make_book(path, shift=0):
    """Mimics TOPICLOG v0.4: a title row, headers on row 2, perp columns far to the right."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2_topic_log"
    ws["A1"] = "TOPIC LOG — decision layer"
    hdr = ["date", "ticker", "demand_score"] + [f"f{i}" for i in range(30 + shift)] + perp_fields.FIELDS
    for c, h in enumerate(hdr, 1):
        ws.cell(2, c, h)
    for r, (d, t) in enumerate([("2026-10-12", "GOLD"), ("2026-10-12", "OIL"),
                                (D, "GOLD"), (D, "OIL")], 3):
        ws.cell(r, 1, d)
        ws.cell(r, 2, t)
    wb.save(path)
    return hdr


def test_header_found_by_name_not_position(tmp_path):
    for shift in (0, 3):
        p = tmp_path / f"b{shift}.xlsx"
        hdr = make_book(p, shift)
        sh = Sheet(open_book(p)["2_topic_log"])
        assert sh.header_row == 2
        assert sh.cols["perp_funding_z"] == hdr.index("perp_funding_z") + 1


def test_upsert_updates_appends_and_guards_order(tmp_path):
    p = tmp_path / "b.xlsx"
    make_book(p)
    sh = Sheet(open_book(p)["2_topic_log"])
    assert sh.upsert("GOLD", D, {"perp_funding_rate": 0.001})[:2] == ("updated", 5)
    act, row, _ = sh.upsert("SILVER", D, {"perp_funding_rate": 0.002})
    assert (act, row) == ("appended", 7)
    with pytest.raises(SheetError, match="older than its latest"):
        sh.upsert("GOLD", dt.date(2026, 10, 1), {})
    with pytest.raises(SheetError, match="no column"):
        sh.upsert("GOLD", D, {"not_a_column": 1})


def test_lock_file_refused(tmp_path):
    p = tmp_path / "b.xlsx"
    make_book(p)
    (tmp_path / "~$b.xlsx").write_text("")
    with pytest.raises(SheetError, match="open in Excel"):
        open_book(p)


def _funding(rates, end):
    return [{"funding_time": (end - dt.timedelta(days=len(rates) - 1 - i)).isoformat(),
             "funding_rate": r, "mark_price": "1"} for i, r in enumerate(rates)]


def test_funding_z_and_no_lookahead():
    asof = dt.datetime(2026, 10, 13, 19, tzinfo=UTC)
    end = dt.datetime(2026, 10, 13, 14, tzinfo=UTC)
    hist = [0.001, -0.001] * 10
    rows = _funding(hist + [0.003], end)
    rows.append({"funding_time": "2026-10-14T14:00:00Z", "funding_rate": 9.0, "mark_price": "1"})  # future
    rate, z, t = perp_fields.funding_fields(rows, asof)
    assert rate == 0.003 and t == end
    assert z == pytest.approx(3.0)


def test_funding_z_blank_when_all_deadband_zeros_or_short():
    end = dt.datetime(2026, 10, 13, 14, tzinfo=UTC)
    assert perp_fields.funding_fields(_funding([0.0] * 25, end), end)[1] is None
    assert perp_fields.funding_fields(_funding([0.001, 0.002] * 3, end), end)[1] is None


def test_oi_change():
    t1 = int(dt.datetime(2026, 10, 13, 18, tzinfo=UTC).timestamp())
    cs = [{"end_period_ts": t1 - 24 * 3600, "open_interest": "1000.00"},
          {"end_period_ts": t1, "open_interest": "1250.00"},
          {"end_period_ts": t1 + 3600, "open_interest": "99999"}]  # after asof
    pct, t = perp_fields.oi_change(cs, dt.datetime(2026, 10, 13, 18, 30, tzinfo=UTC))
    assert pct == pytest.approx(0.25)
    assert perp_fields.oi_change(cs[1:2], dt.datetime(2026, 10, 13, 18, 30, tzinfo=UTC)) == (None, None)


class FakeFetch:
    @staticmethod
    def funding_history(t, s, e):
        return _funding([0.001, -0.001] * 10 + [0.003], dt.datetime(2026, 10, 13, 14, tzinfo=UTC))

    @staticmethod
    def candles(t, s, e, period_minutes=60):
        t1 = int(dt.datetime(2026, 10, 13, 18, tzinfo=UTC).timestamp())
        return [{"end_period_ts": t1 - 86400, "open_interest": "800"}, {"end_period_ts": t1, "open_interest": "1000"}]


def test_fill_perps_end_to_end(tmp_path):
    p = tmp_path / "b.xlsx"
    make_book(p)
    asof = dt.datetime(2026, 10, 13, 19, tzinfo=UTC)
    rep = fill_perps.run(p, D, asof, CFG, fetch=FakeFetch, log=lambda *_: None)
    assert [r[0] for r in rep] == ["GOLD", "OIL"]
    ws = openpyxl.load_workbook(p)["2_topic_log"]
    sh = Sheet(ws)
    g, o = sh.find("GOLD", D), sh.find("OIL", D)
    assert ws.cell(g, sh.cols["perp_funding_rate"]).value == 0.003
    assert ws.cell(g, sh.cols["perp_oi_chg_pct"]).value == pytest.approx(0.25)
    assert ws.cell(o, sh.cols["perp_funding_rate"]).value is None  # no Kalshi perp for OIL
    assert ws.cell(3, sh.cols["perp_funding_rate"]).value is None  # yesterday's row untouched
    assert (tmp_path / "b.bak.xlsx").exists()
