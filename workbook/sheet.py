"""Write Topic Log rows into the Excel workbook by column header, never by column position.

WO-TOPICLOG-ALERTS-01 requires two things of every writer:
  1. find columns by header name, so inserted or moved columns never shift data;
  2. keep each ticker's rows in date order, because RANK reads a ticker's LAST row as its latest.

`Sheet` enforces both. A row is identified by (ticker, date): writing an existing key updates that
row in place; a new key is appended below the last used row, and only if its date is not older than
the ticker's latest row.

Macro-enabled files (.xlsm) are opened with keep_vba so the macro survives a save.
"""
from __future__ import annotations

import datetime as dt
import pathlib
import shutil

import openpyxl


class SheetError(ValueError):
    pass


def as_date(v):
    """Cell value -> date (accepts datetime, date, 'YYYY-MM-DD...' strings); None otherwise."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return dt.date.fromisoformat(str(v).strip()[:10])
    except ValueError:
        return None


def open_book(path):
    path = pathlib.Path(path)
    lock = path.with_name("~$" + path.name)
    if lock.exists():
        raise SheetError(f"{path.name} is open in Excel (lock file {lock.name}); close it first")
    return openpyxl.load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")


def save_book(wb, path, backup=True):
    path = pathlib.Path(path)
    if backup and path.exists():
        shutil.copy2(path, path.with_name(path.stem + ".bak" + path.suffix))
    wb.save(path)


class Sheet:
    def __init__(self, ws, ticker_col="ticker", date_col="date", scan_rows=10):
        self.ws, self.ticker_col, self.date_col = ws, ticker_col, date_col
        self.header_row, self.cols = self._find_header(scan_rows)

    def _find_header(self, scan_rows):
        want = {self.ticker_col, self.date_col}
        for r in range(1, scan_rows + 1):
            names = {}
            for c in range(1, self.ws.max_column + 1):
                v = self.ws.cell(r, c).value
                if isinstance(v, str) and v.strip():
                    names.setdefault(v.strip(), c)
            if want <= names.keys():
                return r, names
        raise SheetError(f"sheet {self.ws.title!r}: no header row with {sorted(want)} in rows 1-{scan_rows}")

    def _last_row(self):
        tc = self.cols[self.ticker_col]
        r = self.ws.max_row
        while r > self.header_row and self.ws.cell(r, tc).value in (None, ""):
            r -= 1
        return r

    def rows(self):
        """Yield (row_number, ticker, date) for every data row."""
        tc, dc = self.cols[self.ticker_col], self.cols[self.date_col]
        for r in range(self.header_row + 1, self._last_row() + 1):
            t = self.ws.cell(r, tc).value
            if t not in (None, ""):
                yield r, str(t).strip(), as_date(self.ws.cell(r, dc).value)

    def find(self, ticker, date):
        hits = [r for r, t, d in self.rows() if t == ticker and d == date]
        return hits[-1] if hits else None

    def latest_date(self, ticker):
        ds = [d for _, t, d in self.rows() if t == ticker and d is not None]
        return max(ds) if ds else None

    def _check_columns(self, values, strict):
        unknown = [k for k in values if k not in self.cols]
        if unknown and strict:
            raise SheetError(f"sheet {self.ws.title!r} has no column(s) {unknown}; "
                             f"add them in the workbook or rename in the writer")
        return unknown

    def upsert(self, ticker, date, values, strict=True, allow_backfill=False):
        """Update the (ticker, date) row, or append it. Returns ('updated'|'appended', row, skipped_keys)."""
        skipped = self._check_columns(values, strict)
        r = self.find(ticker, date)
        action = "updated"
        if r is None:
            latest = self.latest_date(ticker)
            if latest is not None and date < latest and not allow_backfill:
                raise SheetError(f"{ticker} {date}: older than its latest row ({latest}); appending it "
                                 f"would make RANK read stale data")
            r = self._last_row() + 1
            self.ws.cell(r, self.cols[self.ticker_col], ticker)
            self.ws.cell(r, self.cols[self.date_col], date)
            action = "appended"
        for k, v in values.items():
            if k in self.cols and k not in (self.ticker_col, self.date_col):
                self.ws.cell(r, self.cols[k], v)
        return action, r, skipped
